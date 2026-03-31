import datetime
from docx import Document

def replace_word_fields(input_path, output_path):
    # Lấy thời gian thực từ hệ thống
    now = datetime.datetime.now()
    
    # Truyền giá trị vào biến
    current_day = now.strftime("%d")   # Ví dụ: '12'
    current_month = now.strftime("%m") # Ví dụ: '03'
    current_year = now.strftime("%Y")  # Ví dụ: '2026'
    print(now);
    # Mapping từ khóa trong Word -> Giá trị thực tế
    replacements = {
        "[HOTEN]": "",#Lấy từ csdl
        "[DDMMYYSINH]": "", #Lấy từ csdl
        "[DDMMYYSINH]": "",#Lấy từ csdl
        "[DD]": current_day,   # Đã truyền giá trị ngày vào đây
        "[MM]": current_month, # Đã truyền giá trị tháng vào đây
        "[YY]": current_year,   # Đã truyền giá trị năm vào đây
        "[MAIL]": "",#Lấy từ csdl
        "[MSSV]": "",#Lấy từ csdl
        "[LOP]": "",#Lấy từ csdl
        "[KHOA]": "",#Lấy từ csdl
        "[NGANH]": ""#Lấy từ csdl
    }

    doc = Document(input_path)

    # Hàm xử lý thay thế
    for p in doc.paragraphs:
        for key, value in replacements.items():
            if key in p.text:
                for run in p.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, value)

    # Xử lý tương tự cho bảng nếu có
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for key, value in replacements.items():
                        if key in p.text:
                            for run in p.runs:
                                if key in run.text:
                                    run.text = run.text.replace(key, value)

    doc.save(output_path)
    print(f"Đã truyền dữ liệu thành công vào file: {output_path}")

# Chạy lệnh
replace_word_fields('MauTungDon/MauSo5.docx', 'file_ket_qua.docx')